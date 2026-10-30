import os
import io
from functools import wraps
from datetime import datetime
from flask import Flask, render_template, request, jsonify, redirect, url_for, flash, session, send_file
from flask_sqlalchemy import SQLAlchemy
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------------------
# App configuration
# ---------------------------------------------------------------------------
basedir = os.path.abspath(os.path.dirname(__file__))

app = Flask(__name__)
app.config["SECRET_KEY"] = "nutrimais-parceiros-secret-2026"
app.config["SQLALCHEMY_DATABASE_URI"] = "sqlite:///" + os.path.join(basedir, "parceiros.db")
app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False

# Senha padrão de administrador (simples e direta)
ADMIN_PASSWORD = os.environ.get("ADMIN_PASSWORD", "nutrimais2026")

db = SQLAlchemy(app)

@app.context_processor
def inject_admin_status():
    """Make is_admin boolean available in all Jinja templates."""
    return dict(is_admin=session.get("is_admin", False))


def admin_required(f):
    """Decorator to protect admin-only pages."""
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if not session.get("is_admin"):
            return redirect(url_for("index", login_required=1))
        return f(*args, **kwargs)
    return decorated_function

# ---------------------------------------------------------------------------
# Models
# ---------------------------------------------------------------------------

class Nutricionista(db.Model):
    """Registered nutritionists who can refer partners."""
    __tablename__ = "nutricionistas"

    id = db.Column(db.Integer, primary_key=True)
    nome = db.Column(db.String(150), nullable=False)
    ativo = db.Column(db.Boolean, default=True)
    criado_em = db.Column(db.DateTime, default=datetime.utcnow)

    parceiros = db.relationship("Parceiro", backref="nutricionista", lazy=True)

    def to_dict(self):
        return {
            "id": self.id,
            "nome": self.nome,
            "ativo": self.ativo,
            "criado_em": self.criado_em.strftime("%d/%m/%Y") if self.criado_em else "",
        }


class Parceiro(db.Model):
    """Strategic partner registered by a nutritionist."""
    __tablename__ = "parceiros"

    id = db.Column(db.Integer, primary_key=True)
    segmentacao = db.Column(db.String(100), nullable=False)
    nome_empresa = db.Column(db.String(200), nullable=False)
    responsavel_nome = db.Column(db.String(200), nullable=False)
    responsavel_telefone = db.Column(db.String(30), nullable=False)
    responsavel_email = db.Column(db.String(200), nullable=False)
    avaliacao = db.Column(db.Integer, nullable=False)  # 1-5
    nutricionista_id = db.Column(db.Integer, db.ForeignKey("nutricionistas.id"), nullable=False)
    criado_em = db.Column(db.DateTime, default=datetime.utcnow)

    def to_dict(self):
        return {
            "id": self.id,
            "segmentacao": self.segmentacao,
            "nome_empresa": self.nome_empresa,
            "responsavel_nome": self.responsavel_nome,
            "responsavel_telefone": self.responsavel_telefone,
            "responsavel_email": self.responsavel_email,
            "avaliacao": self.avaliacao,
            "nutricionista_id": self.nutricionista_id,
            "nutricionista_nome": self.nutricionista.nome if self.nutricionista else "",
            "criado_em": self.criado_em.strftime("%d/%m/%Y %H:%M") if self.criado_em else "",
        }


class Segmentacao(db.Model):
    """Business segments that can be assigned to partners and toggled visible/hidden."""
    __tablename__ = "segmentacoes"

    id = db.Column(db.Integer, primary_key=True)
    nome = db.Column(db.String(150), nullable=False)
    ativo = db.Column(db.Boolean, default=True)
    criado_em = db.Column(db.DateTime, default=datetime.utcnow)

    def to_dict(self):
        return {
            "id": self.id,
            "nome": self.nome,
            "ativo": self.ativo,
            "criado_em": self.criado_em.strftime("%d/%m/%Y") if self.criado_em else "",
        }


# ---------------------------------------------------------------------------
# Seed data
# ---------------------------------------------------------------------------

def seed_database():
    """Create initial nutritionists and segmentations if tables are empty."""
    if Nutricionista.query.count() == 0:
        nomes = ["Joyce Aliprandi", "Cassiane Oliveira", "Cynthia Papa"]
        for nome in nomes:
            db.session.add(Nutricionista(nome=nome, ativo=True))
        db.session.commit()

    if Segmentacao.query.count() == 0:
        segs_iniciais = ["AVCB", "Dedetização", "Exames Periódicos"]
        for seg in segs_iniciais:
            db.session.add(Segmentacao(nome=seg, ativo=True))
        db.session.commit()


# ---------------------------------------------------------------------------
# Routes – Pages
# ---------------------------------------------------------------------------

@app.route("/")
def index():
    """Landing page with partner registration form."""
    nutricionistas = Nutricionista.query.filter_by(ativo=True).order_by(Nutricionista.nome).all()
    segmentacoes = Segmentacao.query.filter_by(ativo=True).order_by(Segmentacao.nome).all()
    login_required = request.args.get("login_required") == "1"
    return render_template(
        "index.html",
        nutricionistas=nutricionistas,
        segmentacoes=segmentacoes,
        login_required=login_required,
    )


@app.route("/dashboard")
@admin_required
def dashboard():
    """Dashboard with all registered partners (Admin only)."""
    nutricionistas = Nutricionista.query.filter_by(ativo=True).order_by(Nutricionista.nome).all()
    segmentacoes = Segmentacao.query.order_by(Segmentacao.nome).all()
    return render_template("dashboard.html", nutricionistas=nutricionistas, segmentacoes=segmentacoes)


@app.route("/nutricionistas")
@admin_required
def nutricionistas_page():
    """Page to manage nutritionists (Admin only)."""
    nutricionistas = Nutricionista.query.order_by(Nutricionista.nome).all()
    return render_template("nutricionistas.html", nutricionistas=nutricionistas)


@app.route("/segmentacoes")
@admin_required
def segmentacoes_page():
    """Page to manage partner segmentations (Admin only)."""
    segmentacoes = Segmentacao.query.order_by(Segmentacao.nome).all()
    return render_template("segmentacoes.html", segmentacoes=segmentacoes)


# ---------------------------------------------------------------------------
# Routes – Admin Authentication
# ---------------------------------------------------------------------------

@app.route("/api/admin/login", methods=["POST"])
def api_admin_login():
    """Authenticate administrator with password."""
    data = request.get_json() or {}
    senha = (data.get("senha") or "").strip()
    if senha == ADMIN_PASSWORD:
        session["is_admin"] = True
        return jsonify({"success": True, "message": "Acesso concedido como Administrador!"})
    return jsonify({"error": "Senha incorreta. Tente novamente."}), 401


@app.route("/api/admin/logout", methods=["POST", "GET"])
def api_admin_logout():
    """Exit administrator mode back to Nutritionist view."""
    session.pop("is_admin", None)
    if request.method == "GET":
        return redirect(url_for("index"))
    return jsonify({"success": True, "message": "Modo Administrador encerrado."})


# ---------------------------------------------------------------------------
# Routes – API (Partners)
# ---------------------------------------------------------------------------

@app.route("/api/parceiros", methods=["GET"])
def api_get_parceiros():
    """Return all partners, optionally filtered."""
    query = Parceiro.query

    # Filters
    segmentacao = request.args.get("segmentacao")
    nutricionista_id = request.args.get("nutricionista_id")
    avaliacao = request.args.get("avaliacao")
    busca = request.args.get("busca")

    if segmentacao:
        query = query.filter_by(segmentacao=segmentacao)
    if nutricionista_id:
        query = query.filter_by(nutricionista_id=int(nutricionista_id))
    if avaliacao:
        query = query.filter_by(avaliacao=int(avaliacao))
    if busca:
        term = f"%{busca}%"
        query = query.filter(
            db.or_(
                Parceiro.nome_empresa.ilike(term),
                Parceiro.responsavel_nome.ilike(term),
                Parceiro.responsavel_email.ilike(term),
            )
        )

    parceiros = query.order_by(Parceiro.criado_em.desc()).all()
    return jsonify([p.to_dict() for p in parceiros])


@app.route("/api/parceiros/exportar", methods=["GET"])
@admin_required
def api_exportar_excel():
    """Export partners to a styled, segmented, and ranked Excel spreadsheet."""
    query = Parceiro.query

    # Apply filters if provided
    segmentacao = request.args.get("segmentacao")
    nutricionista_id = request.args.get("nutricionista_id")
    avaliacao = request.args.get("avaliacao")
    busca = request.args.get("busca")

    if segmentacao:
        query = query.filter_by(segmentacao=segmentacao)
    if nutricionista_id:
        query = query.filter_by(nutricionista_id=int(nutricionista_id))
    if avaliacao:
        query = query.filter_by(avaliacao=int(avaliacao))
    if busca:
        term = f"%{busca}%"
        query = query.filter(
            db.or_(
                Parceiro.nome_empresa.ilike(term),
                Parceiro.responsavel_nome.ilike(term),
                Parceiro.responsavel_email.ilike(term),
            )
        )

    # Rankeado por Avaliação decrescente (melhores notas primeiro), depois por segmento e empresa
    parceiros = query.order_by(Parceiro.avaliacao.desc(), Parceiro.segmentacao.asc(), Parceiro.nome_empresa.asc()).all()

    wb = openpyxl.Workbook()

    # ── Cores e Estilos NutriMais ─────────────────────────────────
    c_olive_dark = "2A3B25"
    c_olive_main = "3D5636"
    c_zebra = "F4F7F2"
    c_border = "D6DFD2"

    font_title = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    font_sub = Font(name="Calibri", size=10, italic=True, color="FFFFFF")
    font_th = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_body = Font(name="Calibri", size=10)
    font_bold = Font(name="Calibri", size=10, bold=True)
    font_rank = Font(name="Calibri", size=11, bold=True, color=c_olive_main)
    font_star = Font(name="Calibri", size=10, bold=True, color="B8860B")

    fill_title = PatternFill(start_color=c_olive_dark, end_color=c_olive_dark, fill_type="solid")
    fill_sub = PatternFill(start_color=c_olive_main, end_color=c_olive_main, fill_type="solid")
    fill_th = PatternFill(start_color=c_olive_main, end_color=c_olive_main, fill_type="solid")
    fill_seg_banner = PatternFill(start_color="3D5636", end_color="3D5636", fill_type="solid")
    fill_zebra = PatternFill(start_color=c_zebra, end_color=c_zebra, fill_type="solid")
    fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")

    thin_side = Side(border_style="thin", color=c_border)
    border_cell = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    headers = [
        "Posição / Rank",
        "Avaliação",
        "Segmentação",
        "Nome da Empresa",
        "Responsável",
        "Telefone",
        "E-mail",
        "Nutricionista Indicadora",
        "Data de Cadastro",
    ]

    def format_star(n):
        full = "★" * n
        empty = "☆" * (5 - n)
        return f"{full}{empty} ({n}/5)"

    data_geracao = datetime.now().strftime("%d/%m/%Y às %H:%M")

    # ── ABA 1: Ranking Geral ──────────────────────────────────────
    ws1 = wb.active
    ws1.title = "Ranking Geral"
    ws1.views.sheetView[0].showGridLines = True

    # Title Banner
    ws1.merge_cells("A1:I1")
    t1 = ws1["A1"]
    t1.value = "NUTRIMAIS – RELATÓRIO DE PARCEIROS ESTRATÉGICOS (RANKING GERAL)"
    t1.font = font_title
    t1.fill = fill_title
    t1.alignment = align_center
    ws1.row_dimensions[1].height = 36

    ws1.merge_cells("A2:I2")
    s1 = ws1["A2"]
    s1.value = f"Gerado em: {data_geracao} | Total de Parceiros: {len(parceiros)} | Ordenação: Melhores Avaliações (Rankeado)"
    s1.font = font_sub
    s1.fill = fill_sub
    s1.alignment = align_center
    ws1.row_dimensions[2].height = 20

    ws1.row_dimensions[3].height = 10

    # Headers
    ws1.row_dimensions[4].height = 26
    for col_idx, h in enumerate(headers, 1):
        c = ws1.cell(row=4, column=col_idx, value=h)
        c.font = font_th
        c.fill = fill_th
        c.alignment = align_center
        c.border = border_cell

    # Rows
    r_idx = 5
    for rank, p in enumerate(parceiros, 1):
        ws1.row_dimensions[r_idx].height = 22
        fill_curr = fill_zebra if rank % 2 == 0 else fill_white

        ws1.cell(row=r_idx, column=1, value=f"#{rank}").alignment = align_center
        ws1.cell(row=r_idx, column=2, value=format_star(p.avaliacao)).alignment = align_center
        ws1.cell(row=r_idx, column=3, value=p.segmentacao).alignment = align_center
        ws1.cell(row=r_idx, column=4, value=p.nome_empresa).alignment = align_left
        ws1.cell(row=r_idx, column=5, value=p.responsavel_nome).alignment = align_left
        ws1.cell(row=r_idx, column=6, value=p.responsavel_telefone).alignment = align_center
        ws1.cell(row=r_idx, column=7, value=p.responsavel_email).alignment = align_left
        ws1.cell(row=r_idx, column=8, value=p.nutricionista.nome if p.nutricionista else "").alignment = align_center
        ws1.cell(row=r_idx, column=9, value=p.criado_em.strftime("%d/%m/%Y") if p.criado_em else "").alignment = align_center

        for col_idx in range(1, 10):
            cell = ws1.cell(row=r_idx, column=col_idx)
            cell.fill = fill_curr
            cell.border = border_cell
            if col_idx == 1:
                cell.font = font_rank
            elif col_idx == 2:
                cell.font = font_star
            elif col_idx == 4:
                cell.font = font_bold
            else:
                cell.font = font_body

        r_idx += 1

    # Auto-adjust column widths
    for col in ws1.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row > 2 and cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws1.column_dimensions[col_letter].width = max(max_len + 4, 15)

    # ── ABA 2: Segmentados & Rankeados ────────────────────────────
    ws2 = wb.create_sheet(title="Por Segmentação")
    ws2.views.sheetView[0].showGridLines = True

    ws2.merge_cells("A1:I1")
    t2 = ws2["A1"]
    t2.value = "NUTRIMAIS – PARCEIROS POR SEGMENTAÇÃO (RANKEADOS POR AVALIAÇÃO)"
    t2.font = font_title
    t2.fill = fill_title
    t2.alignment = align_center
    ws2.row_dimensions[1].height = 36

    ws2.merge_cells("A2:I2")
    s2 = ws2["A2"]
    s2.value = f"Gerado em: {data_geracao} | Cada segmento organizado e rankeado das melhores avaliações para as menores"
    s2.font = font_sub
    s2.fill = fill_sub
    s2.alignment = align_center
    ws2.row_dimensions[2].height = 20

    # Group partners by segmentation
    segs_dict = {}
    for p in parceiros:
        segs_dict.setdefault(p.segmentacao, []).append(p)

    r_curr = 4
    for seg_nome, lista in sorted(segs_dict.items()):
        lista_ordenada = sorted(lista, key=lambda x: x.avaliacao, reverse=True)
        avg_stars = sum(x.avaliacao for x in lista_ordenada) / len(lista_ordenada)

        # Segment Banner
        ws2.merge_cells(start_row=r_curr, start_column=1, end_row=r_curr, end_column=9)
        sec_cell = ws2.cell(row=r_curr, column=1)
        sec_cell.value = f"📌 SEGMENTAÇÃO: {seg_nome.upper()} ({len(lista_ordenada)} parceiros cadastrados | Média: {avg_stars:.1f} ★)"
        sec_cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        sec_cell.fill = fill_seg_banner
        sec_cell.alignment = align_left
        ws2.row_dimensions[r_curr].height = 26
        r_curr += 1

        # Table Header
        ws2.row_dimensions[r_curr].height = 24
        for col_idx, h in enumerate(headers, 1):
            c = ws2.cell(row=r_curr, column=col_idx, value=h)
            c.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
            c.fill = PatternFill(start_color="4E6B45", end_color="4E6B45", fill_type="solid")
            c.alignment = align_center
            c.border = border_cell
        r_curr += 1

        # Rows within Segment
        for rank_seg, p in enumerate(lista_ordenada, 1):
            ws2.row_dimensions[r_curr].height = 22
            fill_c = fill_zebra if rank_seg % 2 == 0 else fill_white

            ws2.cell(row=r_curr, column=1, value=f"{rank_seg}º no Segmento").alignment = align_center
            ws2.cell(row=r_curr, column=2, value=format_star(p.avaliacao)).alignment = align_center
            ws2.cell(row=r_curr, column=3, value=p.segmentacao).alignment = align_center
            ws2.cell(row=r_curr, column=4, value=p.nome_empresa).alignment = align_left
            ws2.cell(row=r_curr, column=5, value=p.responsavel_nome).alignment = align_left
            ws2.cell(row=r_curr, column=6, value=p.responsavel_telefone).alignment = align_center
            ws2.cell(row=r_curr, column=7, value=p.responsavel_email).alignment = align_left
            ws2.cell(row=r_curr, column=8, value=p.nutricionista.nome if p.nutricionista else "").alignment = align_center
            ws2.cell(row=r_curr, column=9, value=p.criado_em.strftime("%d/%m/%Y") if p.criado_em else "").alignment = align_center

            for c_i in range(1, 10):
                cell = ws2.cell(row=r_curr, column=c_i)
                cell.fill = fill_c
                cell.border = border_cell
                if c_i == 1:
                    cell.font = font_rank
                elif c_i == 2:
                    cell.font = font_star
                elif c_i == 4:
                    cell.font = font_bold
                else:
                    cell.font = font_body

            r_curr += 1

        # Blank row between segments
        r_curr += 1

    for col in ws2.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row > 2 and cell.value and not str(cell.value).startswith("📌"):
                max_len = max(max_len, len(str(cell.value)))
        ws2.column_dimensions[col_letter].width = max(max_len + 4, 15)

    # ── Abas Individuais por Segmento ─────────────────────────────
    for seg_nome, lista in sorted(segs_dict.items()):
        sheet_title = (seg_nome[:28] + "..") if len(seg_nome) > 30 else seg_nome
        ws_seg = wb.create_sheet(title=sheet_title)
        ws_seg.views.sheetView[0].showGridLines = True

        lista_ordenada = sorted(lista, key=lambda x: x.avaliacao, reverse=True)
        avg = sum(x.avaliacao for x in lista_ordenada) / len(lista_ordenada)

        ws_seg.merge_cells("A1:I1")
        t_s = ws_seg["A1"]
        t_s.value = f"NUTRIMAIS – RANKING DE PARCEIROS: {seg_nome.upper()}"
        t_s.font = font_title
        t_s.fill = fill_title
        t_s.alignment = align_center
        ws_seg.row_dimensions[1].height = 36

        ws_seg.merge_cells("A2:I2")
        sub_s = ws_seg["A2"]
        sub_s.value = f"Total: {len(lista_ordenada)} parceiros | Avaliação Média: {avg:.1f} ★ | Exportado em: {data_geracao}"
        sub_s.font = font_sub
        sub_s.fill = fill_sub
        sub_s.alignment = align_center
        ws_seg.row_dimensions[2].height = 20

        # Header
        ws_seg.row_dimensions[4].height = 26
        for col_idx, h in enumerate(headers, 1):
            c = ws_seg.cell(row=4, column=col_idx, value=h)
            c.font = font_th
            c.fill = fill_th
            c.alignment = align_center
            c.border = border_cell

        # Rows
        r_idx = 5
        for rank_s, p in enumerate(lista_ordenada, 1):
            ws_seg.row_dimensions[r_idx].height = 22
            fill_c = fill_zebra if rank_s % 2 == 0 else fill_white

            ws_seg.cell(row=r_idx, column=1, value=f"#{rank_s} no Segmento").alignment = align_center
            ws_seg.cell(row=r_idx, column=2, value=format_star(p.avaliacao)).alignment = align_center
            ws_seg.cell(row=r_idx, column=3, value=p.segmentacao).alignment = align_center
            ws_seg.cell(row=r_idx, column=4, value=p.nome_empresa).alignment = align_left
            ws_seg.cell(row=r_idx, column=5, value=p.responsavel_nome).alignment = align_left
            ws_seg.cell(row=r_idx, column=6, value=p.responsavel_telefone).alignment = align_center
            ws_seg.cell(row=r_idx, column=7, value=p.responsavel_email).alignment = align_left
            ws_seg.cell(row=r_idx, column=8, value=p.nutricionista.nome if p.nutricionista else "").alignment = align_center
            ws_seg.cell(row=r_idx, column=9, value=p.criado_em.strftime("%d/%m/%Y") if p.criado_em else "").alignment = align_center

            for c_i in range(1, 10):
                cell = ws_seg.cell(row=r_idx, column=c_i)
                cell.fill = fill_c
                cell.border = border_cell
                if c_i == 1:
                    cell.font = font_rank
                elif c_i == 2:
                    cell.font = font_star
                elif c_i == 4:
                    cell.font = font_bold
                else:
                    cell.font = font_body

            r_idx += 1

        for col in ws_seg.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.row > 2 and cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            ws_seg.column_dimensions[col_letter].width = max(max_len + 4, 15)

    # Save to BytesIO buffer
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"NutriMais_Parceiros_Ranking_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.route("/api/parceiros", methods=["POST"])
def api_create_parceiro():
    """Create a new partner."""
    data = request.get_json()

    required = ["segmentacao", "nome_empresa", "responsavel_nome",
                 "responsavel_telefone", "responsavel_email",
                 "avaliacao", "nutricionista_id"]

    for field in required:
        if not data.get(field):
            return jsonify({"error": f"Campo obrigatório ausente: {field}"}), 400

    parceiro = Parceiro(
        segmentacao=data["segmentacao"],
        nome_empresa=data["nome_empresa"],
        responsavel_nome=data["responsavel_nome"],
        responsavel_telefone=data["responsavel_telefone"],
        responsavel_email=data["responsavel_email"],
        avaliacao=int(data["avaliacao"]),
        nutricionista_id=int(data["nutricionista_id"]),
    )
    db.session.add(parceiro)
    db.session.commit()
    return jsonify(parceiro.to_dict()), 201


@app.route("/api/parceiros/<int:parceiro_id>", methods=["DELETE"])
def api_delete_parceiro(parceiro_id):
    """Delete a partner."""
    parceiro = Parceiro.query.get_or_404(parceiro_id)
    db.session.delete(parceiro)
    db.session.commit()
    return jsonify({"message": "Parceiro removido com sucesso."}), 200


@app.route("/api/parceiros/<int:parceiro_id>", methods=["PUT"])
def api_update_parceiro(parceiro_id):
    """Update an existing partner."""
    parceiro = Parceiro.query.get_or_404(parceiro_id)
    data = request.get_json()

    if data.get("segmentacao"):
        parceiro.segmentacao = data["segmentacao"]
    if data.get("nome_empresa"):
        parceiro.nome_empresa = data["nome_empresa"]
    if data.get("responsavel_nome"):
        parceiro.responsavel_nome = data["responsavel_nome"]
    if data.get("responsavel_telefone"):
        parceiro.responsavel_telefone = data["responsavel_telefone"]
    if data.get("responsavel_email"):
        parceiro.responsavel_email = data["responsavel_email"]
    if data.get("avaliacao"):
        parceiro.avaliacao = int(data["avaliacao"])
    if data.get("nutricionista_id"):
        parceiro.nutricionista_id = int(data["nutricionista_id"])

    db.session.commit()
    return jsonify(parceiro.to_dict()), 200


# ---------------------------------------------------------------------------
# Routes – API (Nutritionists)
# ---------------------------------------------------------------------------

@app.route("/api/nutricionistas", methods=["GET"])
def api_get_nutricionistas():
    """Return all nutritionists."""
    nutricionistas = Nutricionista.query.order_by(Nutricionista.nome).all()
    return jsonify([n.to_dict() for n in nutricionistas])


@app.route("/api/nutricionistas", methods=["POST"])
def api_create_nutricionista():
    """Create a new nutritionist."""
    data = request.get_json()
    if not data.get("nome"):
        return jsonify({"error": "Nome é obrigatório."}), 400

    nutri = Nutricionista(nome=data["nome"], ativo=True)
    db.session.add(nutri)
    db.session.commit()
    return jsonify(nutri.to_dict()), 201


@app.route("/api/nutricionistas/<int:nutri_id>", methods=["PUT"])
def api_update_nutricionista(nutri_id):
    """Update a nutritionist."""
    nutri = Nutricionista.query.get_or_404(nutri_id)
    data = request.get_json()

    if "nome" in data:
        nutri.nome = data["nome"]
    if "ativo" in data:
        nutri.ativo = bool(data["ativo"])

    db.session.commit()
    return jsonify(nutri.to_dict()), 200


@app.route("/api/nutricionistas/<int:nutri_id>", methods=["DELETE"])
def api_delete_nutricionista(nutri_id):
    """Delete a nutritionist (only if no partners are linked)."""
    nutri = Nutricionista.query.get_or_404(nutri_id)
    if nutri.parceiros:
        return jsonify({"error": "Não é possível excluir nutricionista com parceiros vinculados."}), 400
    db.session.delete(nutri)
    db.session.commit()
    return jsonify({"message": "Nutricionista removida com sucesso."}), 200


# ---------------------------------------------------------------------------
# Routes – API (Segmentações)
# ---------------------------------------------------------------------------

@app.route("/api/segmentacoes", methods=["GET"])
def api_get_segmentacoes():
    """Return all segmentations, optionally only active ones."""
    apenas_ativas = request.args.get("ativas")
    if apenas_ativas in ["true", "1", "True"]:
        segs = Segmentacao.query.filter_by(ativo=True).order_by(Segmentacao.nome).all()
    else:
        segs = Segmentacao.query.order_by(Segmentacao.nome).all()
    return jsonify([s.to_dict() for s in segs])


@app.route("/api/segmentacoes", methods=["POST"])
def api_create_segmentacao():
    """Create a new segmentation."""
    data = request.get_json()
    nome = (data.get("nome") or "").strip()
    if not nome:
        return jsonify({"error": "Nome da segmentação é obrigatório."}), 400

    # Avoid exact duplicates
    existente = Segmentacao.query.filter(db.func.lower(Segmentacao.nome) == nome.lower()).first()
    if existente:
        return jsonify({"error": "Já existe uma segmentação com este nome."}), 400

    seg = Segmentacao(nome=nome, ativo=True)
    db.session.add(seg)
    db.session.commit()
    return jsonify(seg.to_dict()), 201


@app.route("/api/segmentacoes/<int:seg_id>", methods=["PUT"])
def api_update_segmentacao(seg_id):
    """Update segmentation name and/or visibility (ativo/oculto)."""
    seg = Segmentacao.query.get_or_404(seg_id)
    data = request.get_json()

    if "nome" in data:
        novo_nome = data["nome"].strip()
        if not novo_nome:
            return jsonify({"error": "Nome não pode ser vazio."}), 400
        # If renamed, also update historical partners with this old segment name
        Parceiro.query.filter_by(segmentacao=seg.nome).update({"segmentacao": novo_nome})
        seg.nome = novo_nome

    if "ativo" in data:
        seg.ativo = bool(data["ativo"])

    db.session.commit()
    return jsonify(seg.to_dict()), 200


@app.route("/api/segmentacoes/<int:seg_id>", methods=["DELETE"])
def api_delete_segmentacao(seg_id):
    """Delete a segmentation if not linked to existing partners."""
    seg = Segmentacao.query.get_or_404(seg_id)
    em_uso = Parceiro.query.filter_by(segmentacao=seg.nome).count()
    if em_uso > 0:
        return jsonify({
            "error": f"Não é possível excluir: existem {em_uso} parceiro(s) com esta segmentação. Você pode ocultá-la em vez de excluir."
        }), 400

    db.session.delete(seg)
    db.session.commit()
    return jsonify({"message": "Segmentação removida com sucesso."}), 200


# ---------------------------------------------------------------------------
# Startup
# ---------------------------------------------------------------------------

with app.app_context():
    db.create_all()
    seed_database()

if __name__ == "__main__":
    app.run(debug=True, port=5000)
